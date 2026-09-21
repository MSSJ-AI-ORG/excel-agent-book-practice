"""《클로드로 엑셀 천재되기》 실습 채점기.

쓰는 법:
    python grader/check.py <과제번호> <내 파일>
    python grader/check.py 00-01 00-01_점검.xlsx

과제마다 exercises/<번호>/check.json 에 적힌 규칙으로 채점한다.
기대값은 되도록 practice/ 의 원본 자료에서 채점기가 직접 계산한다(check.json 에 숫자를 박지 않는다).

결과 표시
    ✓  맞음
    ✗  틀림 — 쉬운 이유를 같이 적는다
    !  엑셀에서 한 번 열고 저장해야 채점할 수 있음(수식의 계산값이 파일에 없다)
    ?  채점 규칙이 잘못됨 — 과제를 만든 쪽의 문제

끝낼 때 돌려주는 번호: 0 전부 맞음 · 1 틀린 것 있음 · 2 틀린 것은 없지만 다시 저장 필요 · 3 쓰는 법이나 규칙 문제
"""
from __future__ import annotations

import ast
import hashlib
import json
import re
import sys
import warnings
import zipfile
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

PASS, FAIL, RECALC, ERROR = "pass", "fail", "recalc", "error"
MARK = {PASS: "✓", FAIL: "✗", RECALC: "!", ERROR: "?"}
RECALC_HINT = "엑셀에서 이 파일을 한 번 열고 저장한 뒤 다시 채점하세요."

warnings.filterwarnings("ignore", module="openpyxl")


class ConfigError(Exception):
    """check.json 이나 원본 자료가 규칙과 맞지 않는다."""


class NeedsRecalc(Exception):
    """수식은 있는데 파일에 계산값이 없다."""


@dataclass
class Result:
    status: str
    what: str
    reason: str = ""


# ───────────────────────────── 파일 읽기 ─────────────────────────────

def _stale_cache(path: Path) -> bool:
    """엑셀이 «열 때 전부 다시 계산» 하라고 표시된 파일이면 저장된 계산값을 믿지 않는다."""
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("xl/workbook.xml").decode("utf-8", "replace")
    except (KeyError, zipfile.BadZipFile):
        return False
    m = re.search(r"<(?:\w+:)?calcPr\b[^>]*>", xml)
    return bool(m and re.search(r'fullCalcOnLoad="(1|true)"', m.group(0)))


def _empty_string_results(path: Path) -> set:
    """엑셀이 «빈 글자» 결과로 저장한 수식 칸 (시트 이름, 주소) 모음.

    엑셀은 =IF(…,"확인","") 처럼 결과가 "" 인 수식을 <c t="str"><f>…</f><v/></c> 로 저장한다.
    openpyxl 은 이 값을 None 으로 읽어 «계산값 없음» 과 구별하지 못한다. t="str" 이 붙은 칸만 빈 글자로 본다
    (openpyxl 이 저장한 수식 칸은 t 없이 <v></v> 라 여기에 걸리지 않는다).
    """
    import re
    import zipfile
    out = set()
    try:
        with zipfile.ZipFile(path) as z:
            wb = z.read("xl/workbook.xml").decode("utf-8")
            rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
            target = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
            target.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
            for m in re.finditer(r'<sheet\b[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb):
                name, rid = m.group(1), m.group(2)
                t = target.get(rid, "")
                part = t.lstrip("/") if t.startswith("/") else "xl/" + t
                if part not in z.namelist():
                    continue
                xml = z.read(part).decode("utf-8")
                for c in re.finditer(r'<c\b([^>]*)>(.*?)</c>', xml, re.S):
                    attrs, body = c.group(1), c.group(2)
                    if 't="str"' not in attrs or "<f" not in body:
                        continue
                    if re.search(r'<v>[^<]+</v>', body):
                        continue
                    r = re.search(r'\br="([A-Z]+[0-9]+)"', attrs)
                    if r:
                        name_un = name.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
                        out.add((name_un, r.group(1)))
    except Exception:  # 읽지 못하면 예전처럼 «계산값 없음» 으로 둔다
        return set()
    return out


class Book:
    """통합 문서를 수식 쪽과 값 쪽으로 두 번 읽어 둔다."""

    def __init__(self, path: Path, label: str):
        import openpyxl  # 필요할 때만 불러 쓴다

        self.path, self.label = path, label
        try:
            self.wf = openpyxl.load_workbook(path, data_only=False)
            self.wv = openpyxl.load_workbook(path, data_only=True)
        except PermissionError as e:
            raise ConfigError(f"{path.name} 파일을 열 수 없습니다. 엑셀에서 닫고 다시 해 보세요.") from e
        except Exception as e:  # openpyxl 이 못 읽는 파일
            raise ConfigError(f"{path.name} 은(는) 엑셀 파일로 읽히지 않습니다 ({type(e).__name__}).") from e
        self.stale = _stale_cache(path)
        self.empty_str = _empty_string_results(path)

    def sheet(self, name: str):
        if name not in self.wf.sheetnames:
            raise LookupError(f"{self.label} 에 「{name}」 시트가 없습니다.")
        return self.wf[name], self.wv[name]

    def formula(self, sheet: str, addr: str):
        v = self.sheet(sheet)[0][addr].value
        return v if isinstance(v, str) and v.startswith("=") else None

    def raw(self, sheet: str, addr: str):
        """칸에 적힌 것 그대로(수식이면 수식 글자)."""
        return self.sheet(sheet)[0][addr].value

    def value(self, sheet: str, addr: str):
        """칸의 값. 수식 칸인데 계산값이 없으면 NeedsRecalc."""
        if self.formula(sheet, addr) is None:
            return self.raw(sheet, addr)
        v = self.sheet(sheet)[1][addr].value
        if v is None and not self.stale and (sheet, addr.replace("$", "").upper()) in self.empty_str:
            return ""  # 엑셀은 결과가 빈 글자("")인 수식을 <c t="str"><v/></c> 로 저장한다 — openpyxl 은 None 으로 읽는다
        if v is None:
            raise NeedsRecalc(f"{sheet}!{addr} 에 수식은 있는데 계산값이 저장돼 있지 않습니다.")
        if self.stale:
            raise NeedsRecalc(f"{sheet}!{addr} 에 저장된 계산값을 믿을 수 없습니다"
                              "(엑셀이 아닌 프로그램이 저장해 «열 때 다시 계산» 표시가 붙은 파일).")
        return v


_BOOKS: dict[Path, Book] = {}


def load_book(path: Path, label: str) -> Book:
    key = path.resolve()
    if key not in _BOOKS:
        _BOOKS[key] = Book(key, label)
    return _BOOKS[key]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ───────────────────────────── 값 다루기 ─────────────────────────────

def is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def show(v) -> str:
    if is_number(v):
        return f"{v:,.0f}" if float(v).is_integer() else f"{v:,}"
    if v is None or v == "":
        return "(빈칸)"
    return f"「{v}」"


def same(a, b, tol: float = 1e-6) -> bool:
    if is_number(a) and is_number(b):
        return abs(a - b) <= tol * max(1.0, abs(a), abs(b))
    if a in (None, "") and b in (None, ""):
        return True
    return a == b


def normalize(v, how: str):
    if not isinstance(v, str) or how in (None, "none"):
        return v
    if how == "strip":
        return v.strip()
    if how == "remove_spaces":
        return re.sub(r"\s+", "", v)
    raise ConfigError(f"normalize 에 모르는 값 {how!r}")


def round_excel(x, digits=0):
    """엑셀 ROUND 와 같게 — 딱 절반이면 0 에서 먼 쪽으로(파이썬 round 는 짝수 쪽으로 간다)."""
    q = Decimal(1).scaleb(-int(digits))
    return float(Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP))


def cells_in(ws, ref: str):
    """범위 안 칸 주소를 행 순서로 돌려준다."""
    from openpyxl.utils.cell import range_boundaries, get_column_letter

    c1, r1, c2, r2 = range_boundaries(ref)
    return [f"{get_column_letter(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


# ───────────────────────────── 기대값 계산(expect_from) ─────────────────────────────

class Ctx:
    """채점 한 번에 쓰는 것들 — 저장소 뿌리, 제출 파일."""

    def __init__(self, root: Path, submitted: Book | None):
        self.root, self.submitted = root, submitted

    def book(self, spec: dict) -> Book:
        src = spec.get("source")
        if not src:
            raise ConfigError("expect_from 에 source 가 없습니다")
        if src == "@submitted":
            if self.submitted is None:
                raise ConfigError("제출 파일이 없는데 @submitted 를 썼습니다")
            return self.submitted
        path = self.root / src
        if not path.exists():
            raise ConfigError(f"원본 자료 {src} 가 없습니다")
        return load_book(path, src)


def source_value(ctx: Ctx, book: Book, sheet: str, addr: str):
    """원본 자료 칸의 값. 원본(practice) 쪽 수식 칸은 채점기가 계산을 믿지 않으므로 규칙 오류로 본다."""
    if book is not ctx.submitted and book.formula(sheet, addr) is not None:
        raise ConfigError(f"원본 {book.label} {sheet}!{addr} 은 수식 칸입니다 — 입력 칸에서 계산하도록 규칙을 고쳐야 합니다")
    return book.value(sheet, addr)


def _cond_ok(ctx: Ctx, cond: dict, v) -> bool:
    v = normalize(v, cond.get("normalize"))
    if cond.get("not_blank"):
        return v not in (None, "")
    if "equals" in cond:
        want = cond["equals"]
        if isinstance(want, dict):
            want = resolve(ctx, want["from"])
        return v == normalize(want, cond.get("normalize"))
    for key, fn in (("gte", lambda a, b: a >= b), ("lte", lambda a, b: a <= b),
                    ("gt", lambda a, b: a > b), ("lt", lambda a, b: a < b)):
        if key in cond:
            return is_number(v) and fn(v, cond[key])
    raise ConfigError(f"where 조건을 이해하지 못했습니다: {cond}")


def _rows_selected(ctx: Ctx, book: Book, sheet: str, length: int, where: list) -> list[bool]:
    keep = [True] * length
    ws = book.sheet(sheet)[0]
    for cond in where:
        addrs = cells_in(ws, cond["range"])
        if len(addrs) != length:
            raise ConfigError(f"where 범위 {cond['range']} 의 길이가 계산 범위와 다릅니다")
        for i, a in enumerate(addrs):
            keep[i] = keep[i] and _cond_ok(ctx, cond, source_value(ctx, book, sheet, a))
    return keep


def _numbers(vals):
    return [v for v in vals if is_number(v)]


def op_values(ctx, book, spec):
    sheet = spec["sheet"]
    addrs = cells_in(book.sheet(sheet)[0], spec["range"])
    keep = _rows_selected(ctx, book, sheet, len(addrs), spec.get("where", []))
    return [source_value(ctx, book, sheet, a) for a, k in zip(addrs, keep) if k]


def op_sumproduct(ctx, book, spec):
    sheet = spec["sheet"]
    ws = book.sheet(sheet)[0]
    cols = [cells_in(ws, r) for r in spec["ranges"]]
    if len({len(c) for c in cols}) != 1:
        raise ConfigError("sumproduct 범위들의 길이가 다릅니다")
    keep = _rows_selected(ctx, book, sheet, len(cols[0]), spec.get("where", []))
    total = 0
    for i, k in enumerate(keep):
        if not k:
            continue
        vals = [source_value(ctx, book, sheet, col[i]) for col in cols]
        if all(is_number(v) for v in vals):
            p = 1
            for v in vals:
                p *= v
            total += p
    return total


def op_lookup(ctx, book, spec):
    sheet = spec["sheet"]
    ws = book.sheet(sheet)[0]
    keys, outs = cells_in(ws, spec["key_range"]), cells_in(ws, spec["return_range"])
    want = spec["equals"] if not isinstance(spec["equals"], dict) else resolve(ctx, spec["equals"]["from"])
    for k, o in zip(keys, outs):
        if normalize(source_value(ctx, book, sheet, k), spec.get("normalize")) == want:
            return source_value(ctx, book, sheet, o)
    raise ConfigError(f"lookup: {sheet}!{spec['key_range']} 에서 {want!r} 을 찾지 못했습니다")


def op_defined_name(book, spec):
    names = book.wf.defined_names
    if spec["op"] == "defined_name_count":
        return sum(1 for n in names.values() if not getattr(n, "hidden", False))
    dn = names.get(spec["name"])
    if dn is None:
        raise ConfigError(f"이름 정의 {spec['name']} 가 원본에 없습니다")
    text = dn.attr_text
    try:
        return float(text) if "." in text else int(text)
    except ValueError:
        m = re.fullmatch(r"'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)", text)
        if not m:
            raise ConfigError(f"이름 정의 {spec['name']} 의 대상 {text} 를 읽지 못했습니다")
        return book.value(m.group(1), f"{m.group(2)}{m.group(3)}")


def resolve(ctx: Ctx, spec: dict):
    """expect_from 규칙 하나를 계산해 값 하나를 돌려준다."""
    if "calc" in spec:
        return calc(ctx, spec["calc"])
    book, op = ctx.book(spec), spec.get("op", "cell")
    if op == "cell":
        return source_value(ctx, book, spec["sheet"], spec["cell"])
    if op == "sheet_count":
        return len(book.wf.sheetnames)
    if op in ("defined_name_count", "defined_name_value"):
        return op_defined_name(book, spec)
    if op == "sumproduct":
        return op_sumproduct(ctx, book, spec)
    if op == "lookup":
        return op_lookup(ctx, book, spec)
    vals = op_values(ctx, book, spec)
    if op == "sum":
        return sum(_numbers(vals))
    if op == "count":
        return sum(1 for v in vals if v not in (None, ""))
    if op == "count_distinct":
        how = spec.get("normalize")
        return len({normalize(v, how) for v in vals if v not in (None, "")})
    raise ConfigError(f"모르는 op {op!r}")


_CALC_FUNCS = {"round_excel": round_excel, "abs": abs, "min": min, "max": max,
               "if_": lambda c, a, b: a if c else b}
_CALC_NODES = (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Compare, ast.BoolOp, ast.Call, ast.Name,
               ast.Constant, ast.Load, ast.Add, ast.Sub, ast.Mult, ast.Div, ast.USub, ast.Gt, ast.GtE,
               ast.Lt, ast.LtE, ast.Eq, ast.NotEq, ast.And, ast.Or)


def calc(ctx: Ctx, spec: dict):
    """{"expr": "round_excel(a*r*0.5, -3)", "vars": {"a": {...}, "r": {...}}} — 안전한 식만 받는다."""
    tree = ast.parse(spec["expr"], mode="eval")
    for node in ast.walk(tree):
        if not isinstance(node, _CALC_NODES):
            raise ConfigError(f"calc 식에 쓸 수 없는 것: {type(node).__name__}")
        if isinstance(node, ast.Call) and not (isinstance(node.func, ast.Name) and node.func.id in _CALC_FUNCS):
            raise ConfigError("calc 식에 쓸 수 없는 함수")
    env = {k: resolve(ctx, v) for k, v in spec.get("vars", {}).items()}
    env.update(_CALC_FUNCS)
    return eval(compile(tree, "<calc>", "eval"), {"__builtins__": {}}, env)  # noqa: S307 — 위에서 노드를 걸렀다


def expected(ctx: Ctx, rule: dict):
    if "expect" in rule:
        return rule["expect"]
    if "expect_from" in rule:
        spec = rule["expect_from"]
        try:
            return resolve(ctx, spec)
        except LookupError as e:
            if spec.get("source") == "@submitted":
                raise
            raise ConfigError(f"기대값을 계산하지 못했습니다: {e}") from e
    raise ConfigError("expect 나 expect_from 이 없습니다")


def describe_source(rule: dict) -> str:
    spec = rule.get("expect_from")
    if not spec:
        return ""
    note = rule.get("source_note")
    return f" ({note})" if note else " (원본 자료에서 채점기가 직접 계산한 값)"


# ───────────────────────────── 채점 규칙 ─────────────────────────────
# 규칙 하나 = 함수 하나. (ctx, 제출 파일, 규칙) → Result

def _need_sheet(book: Book, name: str):
    if name not in book.wf.sheetnames:
        raise LookupError(f"「{name}」 시트가 없습니다. 시트 이름이 과제와 같은지 보세요.")


def rule_sheet_exists(ctx, book, rule):
    _need_sheet(book, rule["sheet"])
    return PASS, ""


def rule_cell_equals(ctx, book, rule):
    sheet, addr = rule["sheet"], rule["cell"]
    _need_sheet(book, sheet)
    want = expected(ctx, rule)
    got = book.value(sheet, addr)
    if same(got, want, rule.get("tolerance", 1e-6)):
        return PASS, ""
    return FAIL, f"{sheet}!{addr} 에 {show(want)} 이(가) 있어야 하는데 {show(got)} 이(가) 있습니다.{describe_source(rule)}"


def _func_names(formula: str) -> set[str]:
    return {m.upper() for m in re.findall(r"([A-Za-z][A-Za-z0-9.]*)\s*\(", formula)}


def rule_cell_is_formula(ctx, book, rule):
    sheet, addr = rule["sheet"], rule["cell"]
    _need_sheet(book, sheet)
    f = book.formula(sheet, addr)
    if f is None:
        raw = book.raw(sheet, addr)
        return FAIL, f"{sheet}!{addr} 에 수식이 아니라 값 {show(raw)} 이(가) 들어 있습니다. 자료가 바뀌면 따라 바뀌도록 수식으로 넣어야 합니다."
    missing = [fn for fn in rule.get("functions", []) if fn.upper() not in _func_names(f)]
    if missing:
        return FAIL, f"{sheet}!{addr} 수식 {f} 에 {', '.join(missing)} 가 없습니다."
    return PASS, ""


def rule_cell_is_constant(ctx, book, rule):
    sheet, addr = rule["sheet"], rule["cell"]
    _need_sheet(book, sheet)
    f = book.formula(sheet, addr)
    if f is not None:
        return FAIL, f"{sheet}!{addr} 은(는) 손으로 적은 값이어야 하는데 수식 {f} 이(가) 들어 있습니다."
    return PASS, ""


def _norm_formula(f: str) -> str:
    return re.sub(r"\s+", "", f).upper()


def rule_range_same_as_source(ctx, book, rule):
    """제출 파일의 범위가 원본과 같은가 — 값 칸은 값으로, 수식 칸은 수식 글자로 댄다."""
    sheet = rule["sheet"]
    _need_sheet(book, sheet)
    src = ctx.book({"source": rule["source"]})
    src_sheet = rule.get("source_sheet", sheet)
    if src_sheet not in src.wf.sheetnames:
        raise ConfigError(f"원본 {rule['source']} 에 「{src_sheet}」 시트가 없습니다")
    diffs = []
    for addr in cells_in(book.sheet(sheet)[0], rule["range"]):
        sf, mf = src.formula(src_sheet, addr), book.formula(sheet, addr)
        if sf is not None:
            if mf is None or _norm_formula(sf) != _norm_formula(mf):
                diffs.append(f"{addr}(수식 {sf} → {mf or show(book.raw(sheet, addr))})")
        elif mf is not None or not same(src.raw(src_sheet, addr), book.raw(sheet, addr)):
            diffs.append(f"{addr}({show(src.raw(src_sheet, addr))} → {mf or show(book.raw(sheet, addr))})")
    if not diffs:
        return PASS, ""
    shown = ", ".join(diffs[:5]) + (f" 외 {len(diffs) - 5}칸" if len(diffs) > 5 else "")
    return FAIL, f"{sheet} 시트에서 고치면 안 되는 칸이 바뀌었습니다: {shown}"


def _manifest(ctx) -> dict:
    path = ctx.root / "practice" / "MANIFEST.json"
    return json.loads(path.read_text(encoding="utf-8"))["files"]


def rule_practice_untouched(ctx, book, rule):
    """practice 폴더의 원본 파일이 처음 받은 그대로인가 — 지문(sha256)으로 댄다."""
    rel = rule["file"]
    key = rel.split("practice/", 1)[-1]
    info = _manifest(ctx).get(key)
    if info is None:
        raise ConfigError(f"practice/MANIFEST.json 에 {key} 가 없습니다")
    path = ctx.root / rel
    if not path.exists():
        return FAIL, f"원본 {rel} 이(가) 없어졌습니다. 저장소를 다시 받아 practice 폴더를 되살리세요."
    if sha256(path) != info["sha256"]:
        return FAIL, (f"원본 {rel} 이(가) 처음 받은 파일과 다릅니다. 원본을 고친 것입니다. "
                      "저장소에서 원본을 다시 받고, 다음부터는 사본에서 작업하세요.")
    return PASS, ""


def rule_same_bytes_as(ctx, book, rule):
    path = ctx.root / rule["file"]
    if sha256(book.path) == sha256(path):
        return PASS, ""
    return FAIL, f"제출한 파일이 {rule['file']} 와 바이트까지 같아야 하는데 다릅니다."


def rule_font_bold(ctx, book, rule):
    sheet = rule["sheet"]
    _need_sheet(book, sheet)
    ws = book.sheet(sheet)[0]
    bad = [a for a in cells_in(ws, rule["range"]) if not (ws[a].font and ws[a].font.b)]
    return (PASS, "") if not bad else (FAIL, f"{sheet}!{', '.join(bad[:5])} 가 굵게 되어 있지 않습니다.")


def rule_number_format(ctx, book, rule):
    sheet = rule["sheet"]
    _need_sheet(book, sheet)
    ws = book.sheet(sheet)[0]
    bad = [a for a in cells_in(ws, rule["range"]) if ws[a].number_format != rule["format"]]
    if not bad:
        return PASS, ""
    return FAIL, f"{sheet}!{bad[0]} 의 표시 형식이 {ws[bad[0]].number_format} 입니다({rule['format']} 이어야 합니다)."


def rule_fill_present(ctx, book, rule):
    sheet = rule["sheet"]
    _need_sheet(book, sheet)
    ws = book.sheet(sheet)[0]
    bad = [a for a in cells_in(ws, rule["range"]) if not (ws[a].fill and ws[a].fill.fill_type)]
    return (PASS, "") if not bad else (FAIL, f"{sheet}!{', '.join(bad[:5])} 에 채우기 색이 없습니다.")


def rule_freeze_panes(ctx, book, rule):
    sheet = rule["sheet"]
    _need_sheet(book, sheet)
    got = book.sheet(sheet)[0].freeze_panes
    if got == rule["cell"]:
        return PASS, ""
    return FAIL, f"{sheet} 시트의 틀 고정 자리가 {got or '(없음)'} 입니다({rule['cell']} 이어야 합니다)."


def rule_column_width_min(ctx, book, rule):
    sheet = rule["sheet"]
    _need_sheet(book, sheet)
    ws = book.sheet(sheet)[0]
    bad = []
    for col in rule["columns"]:
        dim = ws.column_dimensions.get(col)
        width = dim.width if dim is not None and dim.width else None
        if width is None or width < rule["min"]:
            bad.append(f"{col}열({'기본 너비' if width is None else round(width, 1)})")
    if not bad:
        return PASS, ""
    return FAIL, f"{sheet} 시트 {', '.join(bad)} 이 좁아 숫자가 ##### 로 보일 수 있습니다."


RULES = {
    "sheet_exists": rule_sheet_exists,
    "cell_equals": rule_cell_equals,
    "cell_is_formula": rule_cell_is_formula,
    "cell_is_constant": rule_cell_is_constant,
    "range_same_as_source": rule_range_same_as_source,
    "practice_untouched": rule_practice_untouched,
    "same_bytes_as": rule_same_bytes_as,
    "font_bold": rule_font_bold,
    "number_format": rule_number_format,
    "fill_present": rule_fill_present,
    "freeze_panes": rule_freeze_panes,
    "column_width_min": rule_column_width_min,
}


# ───────────────────────────── 글 파일(.bas·.md) 규칙 ─────────────────────────────

def read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ConfigError(f"{path.name} 을(를) 글자로 읽지 못했습니다")


def _code_lines(text: str, ignore_comments: bool) -> str:
    if not ignore_comments:
        return text
    # 주석 줄은 비워 둔다(지우면 줄 번호가 어긋난다)
    return "\n".join("" if line.lstrip().startswith("'") else line for line in text.splitlines())


def rule_text_contains(ctx, text, rule):
    body = _code_lines(text, rule.get("ignore_comments", False))
    if re.search(rule["pattern"], body, flags=re.M):
        return PASS, ""
    return FAIL, rule.get("fail_hint", f"「{rule['pattern']}」 에 맞는 줄이 없습니다.")


def rule_text_absent(ctx, text, rule):
    body = _code_lines(text, rule.get("ignore_comments", False))
    m = re.search(rule["pattern"], body, flags=re.M)
    if not m:
        return PASS, ""
    line_no = body[: m.start()].count("\n") + 1
    line = body.splitlines()[line_no - 1].strip()
    return FAIL, rule.get("fail_hint", f"없어야 할 줄이 남아 있습니다: {line_no}번째 줄 「{line}」")


TEXT_RULES = {"text_contains": rule_text_contains, "text_absent": rule_text_absent}


# ───────────────────────────── 실행 ─────────────────────────────

def load_exercise(root: Path, ex_id: str) -> dict:
    if not re.fullmatch(r"(\d{2}|P\d)-\d{2}", ex_id):
        raise ConfigError(f"과제 번호 {ex_id!r} 는 00-01 · P1-01 같은 모양이어야 합니다")
    path = root / "exercises" / ex_id / "check.json"
    if not path.exists():
        raise ConfigError(f"과제 {ex_id} 가 없습니다 (exercises/{ex_id}/check.json)")
    data = json.loads(path.read_text(encoding="utf-8"))
    if data.get("mode") not in ("auto", "manual"):
        raise ConfigError("check.json 의 mode 는 auto 나 manual 이어야 합니다")
    return data


def run_rule(ctx: Ctx, target, rule: dict, table: dict) -> Result:
    what = rule.get("what", rule.get("type", "?"))
    fn = table.get(rule.get("type"))
    if fn is None:
        return Result(ERROR, what, f"모르는 규칙 종류 {rule.get('type')!r}")
    try:
        status, reason = fn(ctx, target, rule)
    except NeedsRecalc as e:
        return Result(RECALC, what, f"{e} {RECALC_HINT}")
    except LookupError as e:
        return Result(FAIL, what, str(e).strip("'\""))
    except ConfigError as e:
        return Result(ERROR, what, str(e))
    return Result(status, what, reason)


def grade(ex_id: str, file_path: str | Path | None, root: Path = ROOT) -> tuple[dict, list[Result]]:
    """채점하고 (과제 정보, 결과 목록)을 돌려준다. manual 과제는 결과가 빈 목록이다."""
    _BOOKS.clear()
    ex = load_exercise(root, ex_id)
    if ex["mode"] == "manual":
        return ex, []
    if file_path is None:
        raise ConfigError("채점할 파일을 같이 적어 주세요. 예: python grader/check.py 00-01 내파일.xlsx")
    path = Path(file_path)
    if not path.exists():
        raise ConfigError(f"{path} 파일이 없습니다. 경로를 확인하세요.")
    if ex.get("file_kind", "workbook") == "text":
        target, table, ctx = read_text(path), TEXT_RULES, Ctx(root, None)
    else:
        target = load_book(path, "제출 파일")
        table, ctx = RULES, Ctx(root, target)
    return ex, [run_rule(ctx, target, r, table) for r in ex.get("checks", [])]


def exit_code(results: list[Result]) -> int:
    statuses = {r.status for r in results}
    if ERROR in statuses:
        return 3
    if FAIL in statuses:
        return 1
    return 2 if RECALC in statuses else 0


def report(ex: dict, results: list[Result]) -> str:
    lines = [f"과제 {ex['id']} · {ex.get('title', '')}", ""]
    if ex["mode"] == "manual":
        lines.append("이 과제는 스스로 확인하는 과제입니다. 아래를 하나씩 확인해 보세요.")
        lines += [f"  □ {item}" for item in ex.get("checklist", [])]
        return "\n".join(lines)
    for r in results:
        lines.append(f"  {MARK[r.status]} {r.what}")
        if r.reason:
            lines.append(f"      → {r.reason}")
    n_pass = sum(r.status == PASS for r in results)
    lines += ["", f"맞은 항목 {n_pass} / {len(results)}"]
    tail = {0: "전부 맞았습니다.", 1: "✗ 표시한 항목을 고친 뒤 다시 채점해 보세요.",
            2: RECALC_HINT, 3: "? 표시는 과제 쪽 규칙 문제입니다. 저장소를 최신으로 받아 보세요."}
    lines.append(tail[exit_code(results)])
    if ex.get("checklist"):
        lines += ["", "채점기가 못 보는 것 — 스스로 확인하세요:"] + [f"  □ {item}" for item in ex["checklist"]]
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")  # cp949 창에서도 글자가 깨지지 않게
        except (AttributeError, ValueError):
            pass
    args = sys.argv[1:] if argv is None else argv
    if not args or args[0] in ("-h", "--help"):
        print(__doc__)
        return 3 if not args else 0
    try:
        ex, results = grade(args[0], args[1] if len(args) > 1 else None)
    except ConfigError as e:
        print(f"? {e}")
        return 3
    print(report(ex, results))
    return 0 if ex["mode"] == "manual" else exit_code(results)


if __name__ == "__main__":
    sys.exit(main())
