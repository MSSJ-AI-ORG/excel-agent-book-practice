"""채점기 자기 시험 — python -m pytest grader

세 경우를 반드시 본다.
  1. 정답 파일은 전부 ✓
  2. 오답 파일은 «겨눈 항목만» ✗ (나머지는 ✓ 로 남아야 한다 — 음성 대조)
  3. 계산값이 없는 파일은 ✗ 가 아니라 「엑셀에서 열고 저장」 안내(!)

정답·오답 파일(grader/testdata/00-01)은 엑셀(COM)로 만들어 저장한 것이다.
계산값없음 두 파일은 그 정답 파일을 ① openpyxl 로 다시 저장 ② «열 때 다시 계산» 표시만 붙여 만들었다.
"""
from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

import openpyxl
import pytest

import check

ROOT = Path(__file__).resolve().parent.parent
DATA = Path(__file__).resolve().parent / "testdata" / "00-01"

# 00-01 check.json 의 항목 순서
SHEET, ROWS, IS_FORMULA, TOTAL, SALES_SAME, UNTOUCHED = range(6)


def statuses(ex_id, name, root=ROOT):
    _, results = check.grade(ex_id, DATA / name if isinstance(name, str) else name, root=root)
    return [r.status for r in results], results


# ───────── 1. 정답 ─────────

def test_answer_file_passes_every_item():
    got, results = statuses("00-01", "정답.xlsx")
    assert got == [check.PASS] * 6, [(r.what, r.reason) for r in results]
    assert check.exit_code(results) == 0


# ───────── 2. 오답 — 겨눈 항목만 틀려야 한다 ─────────

@pytest.mark.parametrize("name, should_fail", [
    ("오답_행수.xlsx", {ROWS}),
    ("오답_값박음.xlsx", {IS_FORMULA}),
    ("오답_원본고침.xlsx", {TOTAL, SALES_SAME}),
])
def test_wrong_file_fails_only_targeted_items(name, should_fail):
    got, results = statuses("00-01", name)
    failed = {i for i, s in enumerate(got) if s == check.FAIL}
    assert failed == should_fail, [(r.what, r.status, r.reason) for r in results]
    assert all(s == check.PASS for i, s in enumerate(got) if i not in should_fail)
    assert check.exit_code(results) == 1
    for i in should_fail:
        assert results[i].reason, "틀린 항목에는 쉬운 이유가 붙어야 한다"


def test_wrong_value_reason_shows_both_numbers():
    _, results = statuses("00-01", "오답_원본고침.xlsx")
    reason = results[TOTAL].reason
    assert "14,578,000" in reason and "13,338,000" in reason


# ───────── 3. 계산값 없음 — ✗ 가 아니라 안내 ─────────

@pytest.mark.parametrize("name", ["계산값없음_openpyxl저장.xlsx", "계산값없음_다시계산표시.xlsx"])
def test_missing_cached_value_gives_recalc_notice_not_fail(name):
    got, results = statuses("00-01", name)
    assert got[TOTAL] == check.RECALC
    assert check.FAIL not in got
    assert all(s == check.PASS for i, s in enumerate(got) if i != TOTAL)
    assert check.exit_code(results) == 2
    assert check.RECALC_HINT in results[TOTAL].reason


def test_recalc_notice_does_not_fire_when_value_exists():
    """음성 대조: 계산값이 있는 파일에서는 안내가 뜨지 않는다(정답·오답_원본고침)."""
    for name in ("정답.xlsx", "오답_원본고침.xlsx"):
        got, _ = statuses("00-01", name)
        assert check.RECALC not in got


# ───────── 원본을 건드린 경우 ─────────

@pytest.fixture
def fake_root(tmp_path):
    for sub in ("practice", "exercises"):
        shutil.copytree(ROOT / sub, tmp_path / sub)
    return tmp_path


def test_practice_untouched_catches_modified_original(fake_root):
    target = fake_root / "practice" / "이번 달 매출.xlsx"
    wb = openpyxl.load_workbook(target)
    wb["매출"]["B2"] = "태블릿"
    wb.save(target)
    got, results = statuses("00-01", "정답.xlsx", root=fake_root)
    assert got[UNTOUCHED] == check.FAIL
    assert "원본" in results[UNTOUCHED].reason


def test_practice_untouched_passes_on_fresh_copy(fake_root):
    got, _ = statuses("00-01", "정답.xlsx", root=fake_root)
    assert got[UNTOUCHED] == check.PASS


def test_practice_file_missing(fake_root):
    (fake_root / "practice" / "이번 달 매출.xlsx").unlink()
    got, results = statuses("00-01", "정답.xlsx", root=fake_root)
    assert got[UNTOUCHED] == check.FAIL


# ───────── 기대값 계산기 ─────────

def test_round_excel_goes_away_from_zero_on_half():
    assert check.round_excel(172500, -3) == 173000
    assert round(172500, -3) == 172000  # 파이썬 기본 반올림은 짝수 쪽 — 24.4 의 그 함정
    assert check.round_excel(2.5) == 3
    assert check.round_excel(-2.5) == -3
    assert check.round_excel(79050, -3) == 79000


def _ctx():
    return check.Ctx(ROOT, None)


def _independent_dept_sum(normalize_spaces: bool, dept: str) -> int:
    """채점기 코드와 따로, 대장 B·D 열을 직접 훑어 더한다."""
    ws = openpyxl.load_workbook(ROOT / "practice" / "부서 관리대장.xlsx")["대장"]
    total = 0
    for r in range(4, 20):
        b, d = ws[f"B{r}"].value, ws[f"D{r}"].value
        if b is None:
            continue
        key = "".join(b.split()) if normalize_spaces else b
        if key == dept:
            total += d
    return total


@pytest.mark.parametrize("how, spaces", [("none", False), ("remove_spaces", True)])
def test_sum_with_where_matches_independent_loop(how, spaces):
    spec = {"source": "practice/부서 관리대장.xlsx", "sheet": "대장", "op": "sum", "range": "D4:D19",
            "where": [{"range": "B4:B19", "equals": "영업1팀", "normalize": how}]}
    assert check.resolve(_ctx(), spec) == _independent_dept_sum(spaces, "영업1팀")


def test_count_distinct_and_names():
    ctx = _ctx()
    src = "practice/부서 관리대장.xlsx"
    raw = check.resolve(ctx, {"source": src, "sheet": "대장", "op": "count_distinct", "range": "B4:B19"})
    norm = check.resolve(ctx, {"source": src, "sheet": "대장", "op": "count_distinct", "range": "B4:B19",
                               "normalize": "remove_spaces"})
    assert raw > norm >= 1
    wb = openpyxl.load_workbook(ROOT / "practice" / "부서 관리대장.xlsx")
    assert check.resolve(ctx, {"source": src, "op": "defined_name_count"}) == len(wb.defined_names)
    assert check.resolve(ctx, {"source": src, "op": "sheet_count"}) == len(wb.sheetnames)
    rate = check.resolve(ctx, {"source": src, "op": "defined_name_value", "name": "기준요율"})
    assert rate == wb["기준"]["B8"].value
    bonus = check.resolve(ctx, {"source": src, "op": "defined_name_value", "name": "성과가산"})
    assert bonus == int(wb.defined_names["성과가산"].attr_text)


def test_lookup_and_calc_rebuild_settlement_from_inputs():
    """정산!D3 의 규칙을 입력 칸에서 다시 계산 — 원고 6.1 의 손 검산과 같은 길."""
    src = "practice/부서 관리대장.xlsx"
    sales = {"source": src, "sheet": "대장", "op": "sum", "range": "D4:D19",
             "where": [{"range": "B4:B19", "equals": "영업1팀"}]}
    target = {"source": src, "sheet": "기준", "op": "lookup", "key_range": "A3:A5",
              "return_range": "B3:B5", "equals": "영업1팀"}
    rate = {"source": src, "op": "defined_name_value", "name": "기준요율"}
    spec = {"calc": {"expr": "round_excel(if_(s >= t, s*r, s*r*0.5), -3)",
                     "vars": {"s": sales, "t": target, "r": rate}}}
    s = _independent_dept_sum(False, "영업1팀")
    ws = openpyxl.load_workbook(ROOT / "practice" / "부서 관리대장.xlsx")["기준"]
    t, r = ws["B3"].value, ws["B8"].value
    manual = check.round_excel(s * r if s >= t else s * r * 0.5, -3)
    assert check.resolve(_ctx(), spec) == manual


def test_where_gte_and_not_blank():
    src = "practice/부서 관리대장.xlsx"
    n = check.resolve(_ctx(), {"source": src, "sheet": "대장", "op": "count", "range": "D4:D19",
                               "where": [{"range": "D4:D19", "gte": 3000000}]})
    ws = openpyxl.load_workbook(ROOT / "practice" / "부서 관리대장.xlsx")["대장"]
    assert n == sum(1 for r in range(4, 20) if isinstance(ws[f"D{r}"].value, int) and ws[f"D{r}"].value >= 3000000)


def test_equals_from_submitted_cell(tmp_path):
    path = tmp_path / "s.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "매출"
    wb["매출"]["N1"] = "강민서"
    wb.save(path)
    ctx = check.Ctx(ROOT, check.Book(path, "제출 파일"))
    spec = {"source": "practice/이번 달 매출.xlsx", "sheet": "매출", "op": "sumproduct",
            "ranges": ["C2:C11", "D2:D11"],
            "where": [{"range": "A2:A11", "equals": {"from": {"source": "@submitted", "sheet": "매출", "cell": "N1"}}}]}
    src = openpyxl.load_workbook(ROOT / "practice" / "이번 달 매출.xlsx")["매출"]
    manual = sum(src[f"C{r}"].value * src[f"D{r}"].value for r in range(2, 12) if src[f"A{r}"].value == "강민서")
    assert check.resolve(ctx, spec) == manual


def test_source_formula_cell_is_rule_error():
    """원본의 수식 칸(E열)을 기대값 재료로 쓰면 규칙 오류 — 정답에서 거꾸로 읽지 않게 막는다."""
    spec = {"source": "practice/이번 달 매출.xlsx", "sheet": "매출", "op": "sum", "range": "E2:E11"}
    with pytest.raises(check.ConfigError):
        check.resolve(_ctx(), spec)


@pytest.mark.parametrize("expr", ["__import__('os')", "a.real", "open('x')", "[1,2]"])
def test_calc_rejects_unsafe_expressions(expr):
    with pytest.raises(check.ConfigError):
        check.calc(_ctx(), {"expr": expr, "vars": {}})


def test_bad_specs_are_rule_errors():
    ctx = _ctx()
    src = "practice/부서 관리대장.xlsx"
    bad = [
        {"source": "practice/없는파일.xlsx", "sheet": "대장", "op": "cell", "cell": "A1"},
        {"sheet": "대장", "op": "cell", "cell": "A1"},
        {"source": src, "sheet": "대장", "op": "모르는op", "range": "A1:A2"},
        {"source": src, "sheet": "대장", "op": "sum", "range": "D4:D19", "where": [{"range": "B4:B10", "equals": "x"}]},
        {"source": src, "sheet": "대장", "op": "sum", "range": "D4:D19", "where": [{"range": "B4:B19"}]},
        {"source": src, "sheet": "대장", "op": "sum", "range": "D4:D19",
         "where": [{"range": "B4:B19", "equals": "x", "normalize": "이상한값"}]},
        {"source": src, "sheet": "기준", "op": "lookup", "key_range": "A3:A5", "return_range": "B3:B5", "equals": "없는팀"},
        {"source": src, "op": "defined_name_value", "name": "없는이름"},
        {"source": src, "sheet": "대장", "op": "sumproduct", "ranges": ["D4:D19", "D4:D10"]},
        {"source": "@submitted", "sheet": "대장", "op": "cell", "cell": "A1"},
    ]
    for spec in bad:
        with pytest.raises(check.ConfigError):
            check.resolve(ctx, spec)


# ───────── 서식·글 규칙 ─────────

def _write_exercise(root: Path, ex_id: str, data: dict) -> None:
    d = root / "exercises" / ex_id
    d.mkdir(parents=True, exist_ok=True)
    (d / "check.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


def test_format_rules_pass_and_fail(fake_root, tmp_path):
    from openpyxl.styles import Font, PatternFill

    path = tmp_path / "fmt.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "매출"
    ws["A1"], ws["B1"] = "담당자", "품목"
    ws["A1"].font = Font(bold=True)
    ws["C2"] = 1290000
    ws["C2"].number_format = "#,##0"
    ws["E2"].fill = PatternFill("solid", fgColor="DDEBF7")
    ws.freeze_panes = "A2"
    ws.column_dimensions["D"].width = 12
    wb.save(path)
    checks = [
        {"type": "font_bold", "sheet": "매출", "range": "A1:A1"},
        {"type": "font_bold", "sheet": "매출", "range": "A1:B1"},
        {"type": "number_format", "sheet": "매출", "range": "C2:C2", "format": "#,##0"},
        {"type": "number_format", "sheet": "매출", "range": "C2:C3", "format": "#,##0"},
        {"type": "fill_present", "sheet": "매출", "range": "E2:E2"},
        {"type": "fill_present", "sheet": "매출", "range": "E2:E3"},
        {"type": "freeze_panes", "sheet": "매출", "cell": "A2"},
        {"type": "freeze_panes", "sheet": "매출", "cell": "B2"},
        {"type": "column_width_min", "sheet": "매출", "columns": ["D"], "min": 9},
        {"type": "column_width_min", "sheet": "매출", "columns": ["D", "E"], "min": 9},
        {"type": "cell_is_constant", "sheet": "매출", "cell": "C2"},
        {"type": "sheet_exists", "sheet": "없는시트"},
        {"type": "없는규칙"},
    ]
    _write_exercise(fake_root, "99-01", {"id": "99-01", "mode": "auto", "checks": checks})
    got, results = statuses("99-01", path, root=fake_root)
    P, F, E = check.PASS, check.FAIL, check.ERROR
    assert got == [P, F, P, F, P, F, P, F, P, F, P, F, E], [(r.what, r.reason) for r in results]
    assert check.exit_code(results) == 3


def test_text_rules(fake_root, tmp_path):
    bas = tmp_path / "매크로1_정리.bas"
    bas.write_text("Option Explicit\n' Range(\"A1\").Select 는 주석이라 괜찮다\nSub 매크로1()\n"
                   "    Range(\"A3:F3\").Font.Bold = True\n    Range(\"A1\").Select\nEnd Sub\n", encoding="cp949")
    checks = [
        {"type": "text_contains", "pattern": r"Font\.Bold\s*=\s*True"},
        {"type": "text_absent", "pattern": r"\.Select\b", "ignore_comments": True},
        {"type": "text_contains", "pattern": r"^Option Explicit"},
        {"type": "text_absent", "pattern": r"Selection\."},
    ]
    _write_exercise(fake_root, "99-02", {"id": "99-02", "mode": "auto", "file_kind": "text", "checks": checks})
    got, results = statuses("99-02", bas, root=fake_root)
    assert got == [check.PASS, check.FAIL, check.PASS, check.PASS]
    assert "5번째 줄" in results[1].reason


def test_range_same_as_source_reports_formula_replaced_by_value(fake_root, tmp_path):
    path = tmp_path / "copy.xlsx"
    shutil.copy2(DATA / "정답.xlsx", path)
    wb = openpyxl.load_workbook(path)
    wb["매출"]["E3"] = 1550000  # 수식을 값으로 덮어씀 — 4.1 에서 말한 사고
    wb.save(path)
    checks = [{"type": "range_same_as_source", "sheet": "매출", "range": "A1:E12",
               "source": "practice/이번 달 매출.xlsx"},
              {"type": "range_same_as_source", "sheet": "매출", "range": "A1:E12",
               "source": "practice/이번 달 매출.xlsx", "source_sheet": "없는시트"}]
    _write_exercise(fake_root, "99-03", {"id": "99-03", "mode": "auto", "checks": checks})
    got, results = statuses("99-03", path, root=fake_root)
    assert got == [check.FAIL, check.ERROR]
    assert "E3" in results[0].reason


def test_same_bytes_as(fake_root):
    checks = [{"type": "same_bytes_as", "file": "practice/이번 달 매출.xlsx"}]
    _write_exercise(fake_root, "99-04", {"id": "99-04", "mode": "auto", "checks": checks})
    ok, _ = statuses("99-04", fake_root / "practice" / "이번 달 매출.xlsx", root=fake_root)
    bad, _ = statuses("99-04", DATA / "정답.xlsx", root=fake_root)
    assert ok == [check.PASS] and bad == [check.FAIL]


# ───────── 실행·표시 ─────────

def test_manual_exercise_shows_checklist(fake_root, capsys, monkeypatch):
    _write_exercise(fake_root, "99-05", {"id": "99-05", "title": "혼자 확인", "mode": "manual",
                                         "checklist": ["목록을 적었다", "줄 수가 늘지 않았다"]})
    monkeypatch.setattr(check, "ROOT", fake_root)
    monkeypatch.setattr(check.grade, "__defaults__", (fake_root,))
    assert check.main(["99-05"]) == 0
    out = capsys.readouterr().out
    assert "□ 목록을 적었다" in out and "□ 줄 수가 늘지 않았다" in out


def test_main_exit_codes(capsys):
    assert check.main([]) == 3
    assert check.main(["--help"]) == 0
    assert check.main(["0001"]) == 3
    assert check.main(["98-01", "x.xlsx"]) == 3
    assert check.main(["00-01"]) == 3
    assert check.main(["00-01", "없는파일.xlsx"]) == 3
    assert check.main(["00-01", str(DATA / "정답.xlsx")]) == 0
    assert check.main(["00-01", str(DATA / "오답_행수.xlsx")]) == 1
    assert check.main(["00-01", str(DATA / "계산값없음_openpyxl저장.xlsx")]) == 2
    out = capsys.readouterr().out
    assert "✓" in out and "✗" in out and "!" in out


def test_not_a_workbook_is_reported(tmp_path):
    bogus = tmp_path / "가짜.xlsx"
    bogus.write_text("엑셀 파일이 아니다", encoding="utf-8")
    with pytest.raises(check.ConfigError):
        check.grade("00-01", bogus)


def test_bad_mode_is_rejected(fake_root):
    _write_exercise(fake_root, "99-06", {"id": "99-06", "mode": "알아서"})
    with pytest.raises(check.ConfigError):
        check.grade("99-06", None, root=fake_root)


# ───────── 모든 auto 과제 — 정답 예시는 전부 ✓, 오답 예시는 _verify.txt 에 적힌 항목만 ✗ ─────────

_SUBMIT_EXT = {".xlsx", ".xlsm", ".bas", ".csv", ".md", ".txt"}


def _auto_exercises():
    out = []
    for cj in sorted((ROOT / "exercises").glob("*/check.json")):
        if json.loads(cj.read_text(encoding="utf-8")).get("mode") == "auto":
            out.append(cj.parent.name)
    return out


def _example_file(ex_dir: Path, kind: str, verify: str) -> Path:
    files = sorted(p for p in (ex_dir / kind).rglob("*") if p.is_file() and p.suffix.lower() in _SUBMIT_EXT)
    named = [p for p in files if p.relative_to(ROOT).as_posix() in verify]
    assert named or len(files) == 1, f"{ex_dir.name}/{kind} 에서 채점할 파일을 고를 수 없다: {files}"
    return (named or files)[0]


def _verify_fails(verify: str, rel: str) -> set[str]:
    seg = verify[verify.index(rel):]
    stop = re.search(r"\n(맞은 항목|\(끝)", seg)
    seg = seg[: stop.start()] if stop else seg
    return {m.group(1).strip() for m in re.finditer(r"^\s*✗ (.+)$", seg, flags=re.M)}


@pytest.mark.parametrize("ex_id", _auto_exercises())
def test_every_auto_exercise_solution_and_wrong(ex_id):
    ex_dir = ROOT / "exercises" / ex_id
    verify = (ex_dir / "_verify.txt").read_text(encoding="utf-8")
    sol = _example_file(ex_dir, "solution", verify)
    _, res = check.grade(ex_id, sol)
    assert [r.status for r in res] == [check.PASS] * len(res), [(r.what, r.reason) for r in res if r.status != check.PASS]

    wrong = _example_file(ex_dir, "wrong", verify)
    _, res = check.grade(ex_id, wrong)
    fails = {r.what for r in res if r.status == check.FAIL}
    assert fails, "오답 예시가 아무 항목도 ✗ 로 만들지 못했다"
    assert not [r for r in res if r.status in (check.ERROR, check.RECALC)]
    rel = wrong.relative_to(ROOT).as_posix()
    if rel in verify:
        assert fails == _verify_fails(verify, rel), "✗ 항목이 _verify.txt 기록과 다르다"


def test_report_lists_reason_under_failed_item():
    ex, results = check.grade("00-01", DATA / "오답_행수.xlsx")
    text = check.report(ex, results)
    assert "✗ 점검!B1" in text and "→" in text and "맞은 항목 5 / 6" in text


def test_formula_with_empty_string_result_reads_as_blank():
    """엑셀이 결과 "" 로 저장한 수식 칸은 «계산값 없음» 이 아니라 빈 글자다(06-02 대조 E3)."""
    book = check.Book(ROOT / "exercises" / "06-02" / "solution" / "06-02_대조.xlsx", "정답")
    assert book.formula("대조", "E3") is not None
    assert book.value("대조", "E3") == ""


def test_openpyxl_saved_formula_still_needs_recalc(tmp_path):
    """음성 대조 — openpyxl 이 저장한 수식 칸(계산값 없음)은 여전히 «엑셀에서 열고 저장» 안내가 나와야 한다."""
    src = ROOT / "exercises" / "06-02" / "solution" / "06-02_대조.xlsx"
    dst = tmp_path / "resaved.xlsx"
    wb = openpyxl.load_workbook(src)
    wb.save(dst)
    book = check.Book(dst, "다시 저장")
    with pytest.raises(check.NeedsRecalc):
        book.value("대조", "E3")
