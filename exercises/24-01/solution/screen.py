"""24-01 정답 예시 — 엑셀을 띄우지 않고 「부서 관리대장.xlsx」에서 부서별 매출합계를 읽어 낸다.

쓰는 법(저장소 뿌리에서):
    python -X utf8 exercises/24-01/solution/screen.py "practice/부서 관리대장.xlsx" exercises/24-01/solution/out

세 조건(24.1)
  ① 엑셀을 띄우지 않는다 — xlsx 를 파일로만 읽는다(openpyxl).
  ② 돌 때 AI 를 부르지 않는다 — 네트워크를 쓰는 모듈을 하나도 불러오지 않는다.
  ③ 원본을 고치지 않는다 — 결과는 out 폴더에만 쓰고, 돌기 전과 뒤의 원본 지문(sha256)을 대 본다.

이 파일의 정산 시트에는 계산값이 0 으로 저장돼 있다(엑셀로 한 번도 다시 계산해 저장하지 않은 파일).
그래서 정산 시트의 값을 읽지 않고, 대장의 입력 값에서 지금 수식(SUMIF)과 같은 규칙 —
부서 이름이 정확히 같은 줄만 더한다 — 으로 계산한다.
"""
import csv
import hashlib
import sys
from pathlib import Path

import openpyxl


def sha256(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def main(src, out_dir):
    src, out_dir = Path(src), Path(out_dir)
    before = sha256(src)

    wb = openpyxl.load_workbook(src, read_only=True)
    led = {r: [c.value for c in row] for r, row in enumerate(wb["대장"].iter_rows(min_row=4, max_row=19, max_col=6), start=4)}
    depts = [row[0].value for row in wb["정산"].iter_rows(min_row=3, max_row=5, max_col=1)]
    saved = [row[0].value for row in openpyxl.load_workbook(src, read_only=True, data_only=True)["정산"].iter_rows(min_row=3, max_row=5, min_col=2, max_col=2)]
    wb.close()

    # 비었으면 비었다고 알린다 — 0 으로 채우지 않는다
    empty = [f"대장!{col}{r}" for r, vals in led.items() if vals[1] not in (None, "")
             for col, v in zip("ABCDE", vals[:5]) if v in (None, "")]
    blank_rows = [r for r, vals in led.items() if all(v in (None, "") for v in vals)]
    print("부서가 적힌 줄에서 빈 칸:", empty or "없음")
    print("통째로 빈 줄:", blank_rows)
    print("참고 — 정산 시트 B3:B5 에 저장된 계산값:", saved, "(이 값은 쓰지 않는다)")

    sales = {d: sum(vals[3] for vals in led.values() if vals[1] == d and isinstance(vals[3], (int, float))) for d in depts}

    out_dir.mkdir(parents=True, exist_ok=True)
    with open(out_dir / "매출합계.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["부서", "매출합계"])
        for d, v in sales.items():
            w.writerow([d, v])

    after = sha256(src)
    print("원본 지문 전:", before)
    print("원본 지문 뒤:", after)
    print("원본이 그대로인가:", before == after)
    for d, v in sales.items():
        print(f"{d}: {v:,}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
