"""24-01 오답 예시를 만든다 — 흔한 실수 하나: 부서 이름의 띄어쓰기를 지우고 더한다.

쓰는 법(저장소 뿌리에서):
    python -X utf8 exercises/24-01/wrong/make_wrong.py "practice/부서 관리대장.xlsx" exercises/24-01/wrong/out

정답 프로그램(solution/screen.py)과 다른 곳은 한 줄뿐이다 — 부서 이름을 비교하기 전에 띄어쓰기를 지운다.
「영업 1팀」·「영업1 팀」·「영업 2팀」까지 세어져서 값이 «더 그럴듯하게» 커진다.
하지만 지금 정산 수식(SUMIF)은 이름이 정확히 같은 줄만 더한다. 옮길 때는 지금 수식이 하는 대로 옮긴다(24.3).
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

src, out_dir = Path(sys.argv[1]), Path(sys.argv[2])
wb = openpyxl.load_workbook(src, read_only=True)
led = [[c.value for c in r] for r in wb["대장"].iter_rows(min_row=4, max_row=19, max_col=6)]
depts = [row[0].value for row in wb["정산"].iter_rows(min_row=3, max_row=5, max_col=1)]
wb.close()

squash = lambda s: re.sub(r"\s+", "", s) if isinstance(s, str) else s  # noqa: E731  ← 이 한 줄이 실수
sales = {d: sum(v[3] for v in led if squash(v[1]) == d and isinstance(v[3], (int, float))) for d in depts}

out_dir.mkdir(parents=True, exist_ok=True)
with open(out_dir / "매출합계.csv", "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f, lineterminator="\n")
    w.writerow(["부서", "매출합계"])
    for d, v in sales.items():
        w.writerow([d, v])
print(sales)
