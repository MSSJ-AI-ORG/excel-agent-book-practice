"""00-01 기대값을 practice 원본에서 따로 계산한다(정답 파일은 보지 않는다).

쓰는 법: python -X utf8 exercises/00-01/_expected.py
채점기(check.json 의 expect_from)가 계산하는 값과 같은 것을 채점기 코드를 빌리지 않고 다시 센다.
"""
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
ws = openpyxl.load_workbook(ROOT / "practice" / "이번 달 매출.xlsx")["매출"]

rows = [r for r in range(2, 201) if ws[f"A{r}"].value not in (None, "")]
total = sum(ws[f"C{r}"].value * ws[f"D{r}"].value for r in rows)

print("데이터 행 수(A열에 담당자가 적힌 칸):", len(rows))
print("금액 합계(수량×단가, E열 수식은 안 씀):", f"{total:,}")
