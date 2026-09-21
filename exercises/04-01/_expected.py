"""04-01 기대값을 practice 원본에서 따로 계산하고, check.json 에 박은 값과 대 본다.

쓰는 법: python -X utf8 exercises/04-01/_expected.py

원본의 입력 칸(수식이 없는 칸)에서 A1:E12 의 값을 다시 계산해(E=C×D, E12=합계) «바꾸기 전»을 만들고,
시킨 세 칸(A1·B2·C3)을 바꾼 뒤 같은 계산을 다시 해 «바꾼 뒤»를 만든다. 두 표를 칸마다 대서
바뀐 칸 목록을 얻는다. 원본에서 수식이 있던 칸이면 「따라」, 없던 칸이면 「직접」이다.
정답 파일은 보지 않는다.
"""
import json
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ws = openpyxl.load_workbook(ROOT / "practice" / "이번 달 매출.xlsx")["매출"]

ORDERS = {"A1": "담당", "B2": "태블릿", "C3": 9}  # 책 4.1 의 시킬 말 그대로
ADDRS = [f"{c}{r}" for r in range(1, 13) for c in "ABCDE"]  # 행 순서, 같은 행이면 왼쪽부터


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


inputs = {a: ws[a].value for a in ADDRS if not is_formula(ws[a].value)}
formula_cells = {a for a in ADDRS if is_formula(ws[a].value)}


def evaluate(vals):
    out = dict(vals)
    for r in range(2, 12):
        out[f"E{r}"] = vals[f"C{r}"] * vals[f"D{r}"]
    out["E12"] = sum(out[f"E{r}"] for r in range(2, 12))
    return out


before = evaluate(inputs)
after = evaluate({**inputs, **ORDERS})
changed = [a for a in ADDRS if before.get(a) != after.get(a)]

want = {}
for i, a in enumerate(changed, start=2):
    kind = "따라" if a in formula_cells else "직접"
    want.update({f"A{i}": a, f"B{i}": before[a], f"C{i}": after[a], f"D{i}": kind})
    print(a, before[a], "→", after[a], kind)
want[f"A{len(changed) + 2}"] = ""  # 목록 다음 줄은 비어 있어야 한다
print("바꾼 뒤 매출!E3:", f"{after['E3']:,}", "· E12:", f"{after['E12']:,}")

check = json.loads((HERE / "check.json").read_text(encoding="utf-8"))
bad = 0
for rule in check["checks"]:
    if rule.get("sheet") == "변경목록" and "expect" in rule:
        cell = rule["cell"]
        if cell not in want or want[cell] != rule["expect"]:
            bad += 1
            print("어긋남:", cell, "check.json =", rule["expect"], "원본 계산 =", want.get(cell))
    if rule.get("sheet") == "매출" and "expect" in rule and after[rule["cell"]] != rule["expect"]:
        bad += 1
        print("어긋남: 매출", rule["cell"], rule["expect"])

print("OK — check.json 에 박은 값이 원본에서 계산한 값과 같다" if not bad else f"어긋남 {bad}건")
sys.exit(1 if bad else 0)
