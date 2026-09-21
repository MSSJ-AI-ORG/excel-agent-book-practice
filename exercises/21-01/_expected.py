"""21-01 기대값을 practice 원본에서 따로 계산한다(정답 파일은 보지 않는다).

쓰는 법: python -X utf8 exercises/21-01/_expected.py

검사기 여덟 줄 가운데 일곱 줄은 check.json 의 expect_from 으로 채점기가 원본에서 직접 센다.
8번 줄(일자+부서+금액이 같은 줄이 둘 이상인 줄 수)만은 지금 채점기에 셀 방법이 없어 check.json 에 0 을 적었다.
이 스크립트가 그 0 을 원본에서 따로 센다. 나머지 일곱 줄도 같이 세어 둔다(대조용).
"""
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
ws = openpyxl.load_workbook(ROOT / "practice" / "부서 관리대장.xlsx")["대장"]

blank = lambda v: v in (None, "")  # noqa: E731
rows = [r for r in range(4, 20) if not blank(ws[f"B{r}"].value)]
D = {r: ws[f"D{r}"].value for r in rows}
nums = [v for v in D.values() if isinstance(v, (int, float))]
keys = Counter((ws[f"A{r}"].value, ws[f"B{r}"].value, ws[f"D{r}"].value) for r in rows)
dup = sum(1 for r in rows if keys[(ws[f"A{r}"].value, ws[f"B{r}"].value, ws[f"D{r}"].value)] >= 2)

items = [
    ("부서가 적힌 줄 수", len(rows)),
    ("금액 총합", sum(nums)),
    ("부서는 있는데 금액이 빈 줄 수", sum(1 for r in rows if blank(D[r]))),
    ("금액이 음수인 줄 수", sum(1 for v in nums if v < 0)),
    ("금액이 1억을 넘는 줄 수", sum(1 for v in nums if v > 100_000_000)),
    ("부서 표기 가짓수", len({ws[f"B{r}"].value for r in rows})),
    ("부서는 있는데 담당이 빈 줄 수", sum(1 for r in rows if blank(ws[f"E{r}"].value))),
    ("중복 줄 수(일자+부서+금액이 같은 줄)", dup),
]
for i, (name, v) in enumerate(items, start=2):
    print(f"검사기 {i}행  {name}: {v:,}")

assert dup == 0, "check.json 은 원본의 중복 줄 수를 0 으로 적었다"
print("check.json 에 박은 값(중복 0)이 원본 계산과 맞다.")
