"""01-02 기대값을 practice 원본에서 따로 계산한다(정답 파일은 보지 않는다).

쓰는 법: python -X utf8 exercises/01-02/_expected.py
B3(표 수)·B4(차트 수)는 채점기가 아직 세지 못해 스스로 확인하는 항목으로 돌렸다. 여기서는 참고로 같이 센다.
"""
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
wb = openpyxl.load_workbook(ROOT / "practice" / "부서 관리대장.xlsx")
led, base = wb["대장"], wb["기준"]

names = [n for n, d in wb.defined_names.items() if not getattr(d, "hidden", False)]
tables = sum(len(ws.tables) for ws in wb.worksheets)
charts = sum(len(ws._charts) for ws in wb.worksheets)

dept = {r: led[f"B{r}"].value for r in range(4, 20)}
amount = {r: led[f"D{r}"].value for r in range(4, 20)}
labels = {v for v in dept.values() if v not in (None, "")}

exact = sum(amount[r] for r in dept if dept[r] == "영업1팀")
no_space = sum(amount[r] for r in dept if isinstance(dept[r], str) and re.sub(r"\s+", "", dept[r]) == "영업1팀")
target = next(base[f"B{r}"].value for r in range(3, 6) if base[f"A{r}"].value == "영업1팀")

print("B1 시트 수:", len(wb.sheetnames))
print("B2 이름 정의 수(숨김 제외):", len(names), names)
print("B3 표 수(참고):", tables)
print("B4 차트 수(참고):", charts)
print("B5 부서 표기 가짓수:", len(labels), sorted(labels))
print("B6 「영업1팀」과 정확히 같은 줄의 금액 합:", f"{exact:,}")
print("B7 띄어쓰기를 빼고 비교한 금액 합:", f"{no_space:,}")
print("B8 영업1팀 목표:", f"{target:,}")
