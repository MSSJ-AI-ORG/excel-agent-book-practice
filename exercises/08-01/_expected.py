"""08-01 check.json 에서 「확인 필요」를 기대하는 줄(G6·G9·G13·G18)을 practice 원본에서 따로 골라 대 본다.
기준은 시킬 말 그대로 «금액이 300만 이상».

쓰는 법(저장소 맨 위 폴더에서):  python -X utf8 exercises/08-01/_expected.py
"""
import json
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[2]
ws = openpyxl.load_workbook(ROOT / "practice" / "부서 관리대장.xlsx")["대장"]

rows = [r for r in range(4, 20) if isinstance(ws.cell(r, 4).value, (int, float)) and ws.cell(r, 4).value >= 3_000_000]
print("금액이 300만 이상인 줄:", {r: ws.cell(r, 4).value for r in rows})

check = json.loads(Path(__file__).with_name("check.json").read_text(encoding="utf-8"))
written = sorted(int(c["cell"][1:]) for c in check["checks"] if c.get("expect") == "확인 필요")
print("check.json 에서 「확인 필요」를 기대하는 줄:", written)
ok = written == rows
print("같은가:", ok)
sys.exit(0 if ok else 1)
