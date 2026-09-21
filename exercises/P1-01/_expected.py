"""P1-01 기대값을 practice 원본에서 따로 계산하고, check.json 에 박은 값과 대 본다.

쓰는 법: python -X utf8 exercises/P1-01/_expected.py

1) 「강민서 / 모니터」 줄이 원본의 몇 번째 줄인지 찾는다(check.json 은 그 줄의 C열 주소를 쓴다).
2) 원본 입력 칸에서 A1:E12 의 값을 다시 계산해(E=C×D, E12=합계) 바꾸기 전과 바꾼 뒤를 만들고,
   칸마다 대서 바뀐 칸 목록을 얻는다. 원본에 수식이 있던 칸이면 「따라」, 없던 칸이면 「직접」.
정답 파일은 보지 않는다.
"""
import json
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ws = openpyxl.load_workbook(ROOT / "practice" / "이번 달 매출.xlsx")["매출"]
ADDRS = [f"{c}{r}" for r in range(1, 13) for c in "ABCDE"]


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


hits = [r for r in range(2, 12) if ws[f"A{r}"].value == "강민서" and ws[f"B{r}"].value == "모니터"]
assert len(hits) == 1, f"「강민서 / 모니터」 줄이 하나가 아니다: {hits}"
row = hits[0]
print("「강민서 / 모니터」 줄:", row, "· 원래 수량:", ws[f"C{row}"].value)

inputs = {a: ws[a].value for a in ADDRS if not is_formula(ws[a].value)}
formula_cells = {a for a in ADDRS if is_formula(ws[a].value)}


def evaluate(vals):
    out = dict(vals)
    for r in range(2, 12):
        out[f"E{r}"] = vals[f"C{r}"] * vals[f"D{r}"]
    out["E12"] = sum(out[f"E{r}"] for r in range(2, 12))
    return out


before = evaluate(inputs)
after = evaluate({**inputs, f"C{row}": 9})
changed = [a for a in ADDRS if before.get(a) != after.get(a)]

want_list = {}
for i, a in enumerate(changed, start=2):
    kind = "따라" if a in formula_cells else "직접"
    want_list.update({f"A{i}": a, f"B{i}": before[a], f"C{i}": after[a], f"D{i}": kind})
    print(a, before[a], "→", after[a], kind)
want_list[f"A{len(changed) + 2}"] = ""
want_sales = {f"C{row}": 9}
print("바꾼 뒤 매출!E12:", f"{after['E12']:,}")

check = json.loads((HERE / "check.json").read_text(encoding="utf-8"))
bad = 0
for rule in check["checks"]:
    if "expect" not in rule:
        continue
    table = {"변경목록": want_list, "매출": want_sales}.get(rule.get("sheet"))
    if table is None or table.get(rule["cell"], object()) != rule["expect"]:
        bad += 1
        print("어긋남:", rule.get("sheet"), rule["cell"], "check.json =", rule["expect"])

print("OK — check.json 에 박은 값이 원본에서 계산한 값과 같다" if not bad else f"어긋남 {bad}건")
sys.exit(1 if bad else 0)
