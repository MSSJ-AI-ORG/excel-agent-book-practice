# -*- coding: utf-8 -*-
"""P4-01 에서 H열(수식) 값을 G열과 댈 줄을 원본에서 고른다.

원래 월말정리 코드의 조건 그대로 — 대장 D열 금액이 3,000,000 을 넘는 줄. 그 줄만 H 수식 결과가 「확인」이다.
나머지 줄의 수식 결과는 빈 글자("")인데, 엑셀이 저장한 빈 글자 결과를 채점기가 «계산값 없음»과 가르지 못해
그 줄들은 값 대조에서 뺀다(수식이 있는지는 16줄 모두 본다).

쓰는 법: python -X utf8 exercises/P4-01/_expected.py
"""
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
XLSM = ROOT / "practice" / "부서 관리대장.xlsm"
LIMIT = 3000000          # 원래 코드: If ws.Cells(i, 4).Value > 3000000 Then


def rows_over_limit():
    import openpyxl

    ws = openpyxl.load_workbook(XLSM)["대장"]
    return [r for r in range(4, 20) if isinstance(ws[f"D{r}"].value, (int, float)) and ws[f"D{r}"].value > LIMIT]


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    rows = rows_over_limit()
    print("D열이 3,000,000 을 넘는 줄:", rows)
    check = json.loads((Path(__file__).parent / "check.json").read_text(encoding="utf-8"))
    in_check = [int(c["cell"][1:]) for c in check["checks"]
                if c["type"] == "cell_equals" and c["cell"].startswith("H")]
    print("check.json 에서 H 값을 대는 줄:", in_check, "→ 같은가:", in_check == rows)
