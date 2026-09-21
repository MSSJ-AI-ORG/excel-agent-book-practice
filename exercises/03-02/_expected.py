"""03-02 기대값을 practice 원본에서 따로 계산하고, check.json 에 박은 값과 대 본다.

쓰는 법: python -X utf8 exercises/03-02/_expected.py

채점기에는 아직 「정렬 순서」를 보는 규칙이 없다. 그래서 정렬한 뒤 A·B열과 금액(E열)의 값을
check.json 에 줄마다 박아 두었고, 그 값이 원본을 정렬한 결과인지 이 파일이 확인한다.
정렬 기준: 담당자 오름차순, 같은 담당자 안에서는 금액(수량×단가) 내림차순. 정답 파일은 보지 않는다.
"""
import json
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ws = openpyxl.load_workbook(ROOT / "practice" / "이번 달 매출.xlsx")["매출"]

rows = [(ws[f"A{r}"].value, ws[f"B{r}"].value, ws[f"C{r}"].value * ws[f"D{r}"].value) for r in range(2, 12)]
ordered = sorted(rows, key=lambda t: (t[0], -t[2]))  # 한글 이름은 가나다순 = 유니코드 순

want = {}
for i, (name, item, amount) in enumerate(ordered, start=2):
    want[f"A{i}"], want[f"B{i}"], want[f"E{i}"] = name, item, amount
    print(i, name, item, f"{amount:,}")

picked = [t for t in rows if t[0] == "이현우"]
print("G1 이현우 금액 합:", f"{sum(t[2] for t in picked):,}")
print("G2 전체 금액 합:", f"{sum(t[2] for t in rows):,}")
print("G3 필터로 숨는 줄 수:", len(rows) - len(picked))

check = json.loads((HERE / "check.json").read_text(encoding="utf-8"))
bad = 0
for rule in check["checks"]:
    if rule.get("sheet") == "매출" and "expect" in rule:
        cell = rule["cell"]
        if want.pop(cell, object()) != rule["expect"]:
            bad += 1
            print("어긋남:", cell, "check.json =", rule["expect"])
for cell in want:
    bad += 1
    print("check.json 에 없는 칸:", cell, want[cell])

print("OK — check.json 에 박은 정렬 결과가 원본을 정렬한 값과 같다" if not bad else f"어긋남 {bad}건")
sys.exit(1 if bad else 0)
