"""07-01 check.json 에 박아 둔 글자(대장 B5·B8·B14 의 정리 후 부서 이름)를 practice 원본에서 따로 계산해 대 본다.

쓰는 법(저장소 맨 위 폴더에서):  python -X utf8 exercises/07-01/_expected.py
정답 파일은 보지 않는다. 원본 「부서 관리대장.xlsx」만 읽는다.
"""
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[2]
ws = openpyxl.load_workbook(ROOT / "practice" / "부서 관리대장.xlsx")["대장"]

# 공백이 든 부서 칸만 골라, 공백을 모두 뺀 글자를 계산한다
got = {}
for r in range(4, 20):
    v = ws.cell(r, 2).value
    if isinstance(v, str) and re.search(r"\s", v):
        got[f"B{r}"] = re.sub(r"\s+", "", v)
print("공백이 든 칸과 뺀 뒤 글자:", got)
print("공백을 뺀 뒤 표기 수:", len({re.sub(r"\s+", "", ws.cell(r, 2).value) for r in range(4, 20) if ws.cell(r, 2).value}))

check = json.loads(Path(__file__).with_name("check.json").read_text(encoding="utf-8"))
written = {c["cell"]: c["expect"] for c in check["checks"] if c.get("sheet") == "대장" and "expect" in c}
print("check.json 에 적힌 값:", written)
bad = {k: (v, got.get(k)) for k, v in written.items() if got.get(k) != v}
bad.update({k: ("(check.json 에 없음)", v) for k, v in got.items() if k not in written})
print("다른 것:", bad or "없음")
sys.exit(1 if bad else 0)
