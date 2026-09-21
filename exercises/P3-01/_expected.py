"""P3-01 check.json 이 집계!A2:A5 의 기대값으로 가리키는 원본 칸(A2·A3·A5·A8)이
정말 «담당자 이름이 처음 나오는 칸»인지, 그리고 담당자별 금액·건수를 원본에서 따로 계산해 보여 준다.

쓰는 법(저장소 맨 위 폴더에서):  python -X utf8 exercises/P3-01/_expected.py
"""
import json
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[2]
ws = openpyxl.load_workbook(ROOT / "practice" / "이번 달 매출.xlsx")["매출"]

first, amount, count = {}, {}, {}
for r in range(2, 12):
    name, qty, price = ws.cell(r, 1).value, ws.cell(r, 3).value, ws.cell(r, 4).value
    first.setdefault(name, f"A{r}")
    amount[name] = amount.get(name, 0) + qty * price
    count[name] = count.get(name, 0) + 1
print("처음 나오는 칸:", first)
print("담당자별 금액:", amount, "합계", sum(amount.values()))
print("담당자별 건수:", count, "합계", sum(count.values()))

check = json.loads(Path(__file__).with_name("check.json").read_text(encoding="utf-8"))
written = [c["expect_from"]["cell"] for c in check["checks"]
           if c.get("sheet") == "집계" and c.get("cell") in ("A2", "A3", "A4", "A5")]
ok = written == list(first.values())
print("check.json 이 가리키는 칸:", written, "→", "같다" if ok else "다르다")
sys.exit(0 if ok else 1)
