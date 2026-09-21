"""P2-01 check.json 에 박아 둔 값을 practice 원본에서 따로 계산해 대 본다.

- 정리후보 A2:A5: 수식이 하나도 없고, 다른 시트의 수식(과 이름 정의)에 시트 이름이 나오지 않는 시트 — 시트 탭 차례
- 대장 B5·B8·B14: 공백을 뺀 부서 이름
- 대장 G열 「확인 필요」 줄: 금액이 300만 이상인 줄

쓰는 법(저장소 맨 위 폴더에서):  python -X utf8 exercises/P2-01/_expected.py
"""
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[2]
wb = openpyxl.load_workbook(ROOT / "practice" / "부서 관리대장.xlsx")


def formulas(ws):
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]


texts = [f for ws in wb.worksheets for f in formulas(ws)] + [d.attr_text for _, d in wb.defined_names.items()]
cands = [ws.title for ws in wb.worksheets
         if not formulas(ws) and not any(re.search(rf"'?{re.escape(ws.title)}'?!", t) for t in texts)]
led = wb["대장"]
norm = {f"B{r}": re.sub(r"\s+", "", led.cell(r, 2).value) for r in range(4, 20)
        if isinstance(led.cell(r, 2).value, str) and re.search(r"\s", led.cell(r, 2).value)}
flags = [r for r in range(4, 20) if isinstance(led.cell(r, 4).value, (int, float)) and led.cell(r, 4).value >= 3_000_000]

check = json.loads(Path(__file__).with_name("check.json").read_text(encoding="utf-8"))
w_cands = [c["expect"] for c in check["checks"] if c.get("sheet") == "정리후보" and c.get("cell", "").startswith("A")]
w_norm = {c["cell"]: c["expect"] for c in check["checks"] if c.get("sheet") == "대장" and "expect" in c and c.get("cell", "").startswith("B")}
w_flags = sorted(int(c["cell"][1:]) for c in check["checks"] if c.get("expect") == "확인 필요")

result = {"정리후보": (cands, w_cands), "공백 뺀 부서": (norm, w_norm), "확인 필요 줄": (flags, w_flags)}
bad = False
for k, (mine, written) in result.items():
    same = mine == written
    bad |= not same
    print(f"{k}: 원본에서 계산 {mine} / check.json {written} → {'같다' if same else '다르다'}")
sys.exit(1 if bad else 0)
