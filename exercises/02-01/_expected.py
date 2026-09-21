"""02-01 기대값을 practice 원본에서 따로 계산한다(정답 파일은 보지 않는다).

쓰는 법: python -X utf8 exercises/02-01/_expected.py

채점기에는 아직 「값이 있는 마지막 줄 번호」와 「수식 칸 수」를 바로 세는 기능이 없다.
그래서 check.json 은 원본 A열의 데이터 행 수에서 계산한다(행 = 데이터 + 머리글 1 + 합계 1,
수식 칸 = 데이터 행마다 금액 1칸 + 합계 1칸). 이 파일은 그 계산이 원본을 직접 센 값과 같은지 대 본다.
"""
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
wb = openpyxl.load_workbook(ROOT / "practice" / "이번 달 매출.xlsx")
ws = wb["매출"]


def filled(v):
    return v not in (None, "")


cells = [c for row in ws.iter_rows() for c in row if filled(c.value)]
last_row = max(c.row for c in cells)
formulas = [c.coordinate for c in cells if isinstance(c.value, str) and c.value.startswith("=")]
data_rows = sum(1 for r in range(2, 201) if filled(ws[f"A{r}"].value))

print("시트 수:", len(wb.sheetnames))
print("값이 있는 마지막 줄 번호(직접 셈):", last_row)
print("수식 칸 수(직접 셈):", len(formulas), formulas)
print("check.json 식 — 행:", data_rows + 2, "· 수식 칸:", data_rows + 1)

assert last_row == data_rows + 2, "행 수 계산식이 원본과 어긋난다"
assert len(formulas) == data_rows + 1, "수식 칸 계산식이 원본과 어긋난다"
print("OK — check.json 의 계산식이 원본을 직접 센 값과 같다")
