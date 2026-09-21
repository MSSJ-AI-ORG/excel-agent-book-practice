"""19-01 기대값을 practice 원본에서 따로 계산한다(정답 파일은 보지 않는다).

쓰는 법: python -X utf8 exercises/19-01/_expected.py

check.json 은 「2026년 1월 줄 = 원본 대장 4·5·6행」, 「구분이 특판인 줄 = 0건」을 전제로 칸 주소를 적었다.
이 스크립트는 원본 대장의 일자 글자(2026-01-08 · 2026.01.22 · 2026년 2월 14일 세 가지 모양)를 날짜로 읽어
그 전제가 맞는지 다시 센다. 전제가 어긋나면 AssertionError 로 멈춘다.
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
ws = openpyxl.load_workbook(ROOT / "practice" / "부서 관리대장.xlsx")["대장"]


def ym(text):
    if not isinstance(text, str):
        return None
    m = re.fullmatch(r"\s*(\d{4})\s*[-.]\s*(\d{1,2})\s*[-.]\s*(\d{1,2})\s*", text) or \
        re.fullmatch(r"\s*(\d{4})\s*년\s*(\d{1,2})\s*월\s*(\d{1,2})\s*일\s*", text)
    return (int(m.group(1)), int(m.group(2))) if m else None


rows = [r for r in range(4, 20) if ws[f"B{r}"].value not in (None, "")]
jan = [r for r in rows if ym(ws[f"A{r}"].value) == (2026, 1)]
unread = [r for r in rows if ym(ws[f"A{r}"].value) is None]
special_kind = [r for r in rows if ws[f"C{r}"].value == "특판"]
special_note = [r for r in rows if ws[f"F{r}"].value == "특판"]
last = max(rows)

print("자료 줄:", len(rows), "줄 (4~19행 중 부서가 적힌 줄)")
print("날짜로 못 읽은 줄:", unread)
print("2026년 1월 줄:", jan, "→ 금액", [ws[f"D{r}"].value for r in jan], "합", sum(ws[f"D{r}"].value for r in jan))
print("구분(C)이 「특판」인 줄:", special_kind, "(0건이면 멈춰야 한다)")
print("비고(F)가 「특판」인 줄:", special_note, "(구분이 아니다 — 지우면 안 된다)")
print("원본 마지막 자료 줄:", last)

# check.json 이 전제한 칸 주소와 대 본다
assert unread == [], "날짜로 못 읽은 일자가 있다"
assert jan == [4, 5, 6], "check.json 은 1월 줄을 4·5·6행으로 적었다"
assert special_kind == [], "check.json 은 구분이 특판인 줄이 0건이라고 전제했다"
assert special_note == [9], "check.json 은 비고가 특판인 줄이 9행(지운 뒤 6행)이라고 전제했다"
assert last == 19, "check.json 은 마지막 줄 19행이 지운 뒤 16행으로 온다고 전제했다"
shift = len(jan)
cj = json.loads((Path(__file__).parent / "check.json").read_text(encoding="utf-8"))
cells = {(c.get("sheet"), c.get("cell")): c for c in cj["checks"] if c.get("type") == "cell_equals"}
assert cells[("대장", "D6")]["expect_from"]["cell"] == f"D{9}" and 9 - shift == 6
assert cells[("대장", "A16")]["expect_from"]["cell"] == f"A{last}" and last - shift == 16
print("check.json 의 칸 주소가 원본 계산과 맞다.")
