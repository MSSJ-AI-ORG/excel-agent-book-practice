# -*- coding: utf-8 -*-
"""15-01 에서 줄 번호를 박아야 하는 채점 항목의 줄 번호를 원본에서 센다.

합친 시트는 과제대로 머리글 1줄 + 영업1팀 → 영업2팀 → 영업3팀 파일 순서로 이어 붙인다.
그러면 영업3팀 파일의 첫 줄·마지막 줄이 합친 시트 몇 행에 오는지, 데이터가 끝난 다음 행이 몇 행인지가 정해진다.
금액·담당 값 자체는 check.json 이 원본 칸을 직접 가리키므로(expect_from) 여기서 숫자를 박지 않는다.

쓰는 법: python -X utf8 exercises/15-01/_expected.py
"""
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
FOLDER = ROOT / "practice" / "받은파일"


def data_rows(path):
    import openpyxl

    ws = openpyxl.load_workbook(path)["대장"]
    return [r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value not in (None, "")]


def layout():
    files = sorted(FOLDER.glob("*.xlsx"))
    counts = {f.name: len(data_rows(f)) for f in files}
    n1, n2, n3 = (counts[f.name] for f in files)
    return {
        "파일별 데이터 줄 수": counts,
        "영업3팀 첫 줄 → 합친 시트 행": 2 + n1 + n2,
        "영업3팀 마지막 줄 → 합친 시트 행": 1 + n1 + n2 + n3,
        "영업3팀 원본 첫 줄·마지막 줄": [data_rows(files[2])[0], data_rows(files[2])[-1]],
        "데이터가 끝난 다음 행": 2 + n1 + n2 + n3,
    }


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    lay = layout()
    for k, v in lay.items():
        print(f"{k}: {v}")
    check = json.loads((Path(__file__).parent / "check.json").read_text(encoding="utf-8"))
    cells = {c.get("cell") for c in check["checks"] if c.get("sheet") == "합친자료"}
    want = {f"D{lay['영업3팀 첫 줄 → 합친 시트 행']}", f"D{lay['영업3팀 마지막 줄 → 합친 시트 행']}",
            f"A{lay['데이터가 끝난 다음 행']}"}
    print("check.json 이 이 행들을 보고 있나:", want <= cells)
